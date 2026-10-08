from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.throttling import AnonRateThrottle
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework.filters import SearchFilter

from core.permissions import RBACPermission, _has_access
from core.mixins import InstitutionScopedMixin
from .models import DocumentOfficiel, RegistreDiplome
from .serializers import (
    DocumentOfficielSerializer, RegistreDiplomeSerializer, DocumentVerificationSerializer,
)
from core.telechargement import entete_piece_jointe


class VerificationThrottle(AnonRateThrottle):
    """Limite le débit de l'endpoint public /verifier (scope 'verify') pour
    empêcher le balayage de tokens. Par IP (utilisateurs anonymes)."""
    scope = 'verify'


# Mapping type_document -> code module RBAC granulaire (Phase 1 RBAC granulaire)
DOC_TYPE_TO_MODULE = {
    'attestation_inscription': 'doc_attestation',
    'attestation_reussite':    'doc_attestation',
    'releve_complet':          'doc_releve',
    'releve_semestre':         'doc_releve',
    'attestation_diplome':     'doc_diplome',
    'diplome':                 'doc_diplome',
}


def _check_doc_module(user, type_document, action='modifier'):
    """Verifie que l'user a le droit RBAC sur ce type_document.
    Admin/superuser bypassent. Sinon raise PermissionDenied si pas autorise.
    """
    if user.role == 'admin' or user.is_superuser:
        return
    code = DOC_TYPE_TO_MODULE.get(type_document)
    if not code:
        raise PermissionDenied(f'Type de document inconnu : {type_document}')
    if not _has_access(user, code, action):
        raise PermissionDenied(f'Vous n\'avez pas le droit "{action}" sur ce type de document.')


def _serialize_attestation_travail(att):
    """Payload de vérification publique d'une attestation enseignant — MÊME forme
    que le payload étudiant (+ champs prof), consommé par QrVerifyCard côté front."""
    prof = att.prof
    if att.date_debut and att.date_fin:
        periode = f"du {att.date_debut.strftime('%d/%m/%Y')} au {att.date_fin.strftime('%d/%m/%Y')}"
    elif att.annee_universitaire:
        periode = f"Année {att.annee_universitaire}"
    else:
        periode = None
    qualite = {
        'vacataire':   'Enseignant vacataire',
        'permanent':   'Enseignant permanent',
        'contractuel': 'Enseignant contractuel',
        'militaire':   'Enseignant militaire',
    }.get((prof.type or '').lower(), 'Enseignant')
    return {
        'est_valide':          att.est_valide,
        'type_document':       'attestation_travail',
        'type_libelle':        att.titre_document or "Attestation d'enseignement",
        'numero_serie':        att.numero,
        'date_generation':     att.date_generation.isoformat() if att.date_generation else None,
        'annee_universitaire': att.annee_universitaire,
        'titulaire_nom':       prof.nom,
        'etudiant_nom':        prof.nom,
        'etudiant_matricule':  None,
        'photo_url':           None,
        'nni':                 str(prof.NNI),
        'qualite':             qualite,
        'periode':             periode,
        'heures_eq_cm':        f'{att.heures_eq_cm:.2f}',
    }


class DocumentOfficielViewSet(InstitutionScopedMixin, viewsets.ReadOnlyModelViewSet):
    queryset = DocumentOfficiel.objects.select_related('etudiant', 'genere_par').all()
    serializer_class   = DocumentOfficielSerializer
    permission_classes = [RBACPermission]
    filter_backends    = [DjangoFilterBackend, SearchFilter]
    filterset_fields   = ['etudiant', 'type_document', 'est_valide']
    search_fields      = ['numero_serie', 'etudiant__nom', 'etudiant__matricule']
    # Le list/retrieve ne distingue pas les types : on gate sur le registre
    # (le moins privilegie). Les actions custom (generer/regenerer/telecharger)
    # appliquent une verification fine selon doc.type_document via _check_doc_module.
    required_module    = 'doc_registre'

    @action(detail=False, methods=['post'], permission_classes=[RBACPermission])
    def generer(self, request):
        """Génère un nouveau document officiel pour un étudiant.
        Verifie que l'user a le droit RBAC sur le type_document demande."""
        type_doc = (request.data or {}).get('type_document', '')
        _check_doc_module(request.user, type_doc, action='modifier')
        from .services import generer_document
        try:
            data = generer_document(request.data, request.user)
            return Response(data, status=status.HTTP_201_CREATED)
        except Exception as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'], url_path='generer-groupe')
    def generer_groupe(self, request):
        """Génération GROUPÉE : 1 document officiel par étudiant concerné, fusionnés
        en UN seul PDF (renvoyé directement). Filtres : type_document +
        annee_universitaire + filiere (+ semestre pour les relevés, + niveau pour
        cibler une seule promotion quand la filière en porte plusieurs)."""
        d           = request.data or {}
        type_doc    = d.get('type_document', '')
        annee       = d.get('annee_universitaire')
        filiere_id  = d.get('filiere')
        semestre_id = d.get('semestre') or None
        niveau      = d.get('niveau') or None

        _check_doc_module(request.user, type_doc, action='modifier')

        ALLOWED = ('releve_semestre', 'attestation_reussite', 'attestation_inscription', 'attestation_diplome')
        if type_doc not in ALLOWED:
            return Response({'detail': "Type non supporté en génération groupée."},
                            status=status.HTTP_400_BAD_REQUEST)
        if not annee or not filiere_id:
            return Response({'detail': "annee_universitaire et filiere sont requis."},
                            status=status.HTTP_400_BAD_REQUEST)
        if niveau is not None:
            try:
                niveau = int(niveau)
            except (TypeError, ValueError):
                return Response({'detail': "niveau doit être un entier (1, 2, 3…)."},
                                status=status.HTTP_400_BAD_REQUEST)
            if niveau < 1:
                return Response({'detail': "niveau doit être ≥ 1."},
                                status=status.HTTP_400_BAD_REQUEST)

        from .services import generer_documents_groupe
        try:
            pdf_bytes, nb_ok, nb_total, erreurs = generer_documents_groupe(
                type_doc, annee, int(filiere_id),
                int(semestre_id) if semestre_id else None, request.user,
                niveau=niveau,
            )
        except ValueError as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        from django.http import HttpResponse
        # Nom : type_semestre_filiere_niveau_annee
        # (ex. attestation_inscription_LPSEA_L2_2026_2027.pdf). Le niveau figure dans
        # le nom : deux promotions d'une meme filiere produiraient sinon deux fichiers
        # homonymes qui s'ecrasent au telechargement.
        from apps.scolarite.models import Filiere
        from apps.parametres.models import Semestre
        _fil = Filiere.objects.filter(pk=filiere_id).first()
        fil_code = _fil.code if _fil else str(filiere_id)
        sem_code = ''
        if semestre_id:
            _sem = Semestre.objects.filter(pk=semestre_id).first()
            sem_code = _sem.code_semestre if _sem else str(semestre_id)
        niv_code = f'L{niveau}' if niveau else ''
        fname = '_'.join(p for p in (type_doc, sem_code, fil_code, niv_code,
                                     (annee or '').replace('-', '_')) if p) + '.pdf'
        resp = HttpResponse(pdf_bytes, content_type='application/pdf')
        resp['Content-Disposition'] = entete_piece_jointe(fname, inline=True)
        resp['X-Generated'] = str(nb_ok)
        resp['X-Total']     = str(nb_total)
        resp['Access-Control-Expose-Headers'] = 'X-Generated, X-Total'
        return resp

    @action(detail=False, methods=['get', 'post'], url_path='releves-etudiant')
    def releves_etudiant(self, request):
        """Les relevés de notes d'un ou de plusieurs étudiants, en un PDF.

        GET  ?etudiant=ID          : ses semestres, avec moyenne et décision (aperçu).
        POST {etudiants: [ID, ...]} (ou {etudiant: ID}) : leurs relevés (semestres
             qui ont des résultats), étudiant par étudiant dans l'ordre donné,
             fusionnés. Voir releves_etudiant.py.
        """
        from apps.absence.models import Etudiant
        from .releves_etudiant import generer_releves_etudiants, semestres_de

        source = request.query_params if request.method == 'GET' else (request.data or {})
        _check_doc_module(request.user, 'releve_semestre',
                          action='voir' if request.method == 'GET' else 'modifier')
        ids = source.get('etudiants') if request.method == 'POST' else None
        if not ids:
            ids = [source.get('etudiant')]
        try:
            ids = [int(i) for i in ids]
        except (TypeError, ValueError):
            return Response({'detail': 'Étudiant introuvable.'}, status=status.HTTP_400_BAD_REQUEST)
        trouves = Etudiant.objects.in_bulk(ids)
        if not ids or any(i not in trouves for i in ids):
            return Response({'detail': 'Étudiant introuvable.'}, status=status.HTTP_400_BAD_REQUEST)
        etudiants = [trouves[i] for i in dict.fromkeys(ids)]      # ordre donné, sans doublon

        if request.method == 'GET':
            etudiant = etudiants[0]
            return Response({'etudiant': {'id': etudiant.pk, 'matricule': etudiant.matricule,
                                          'nom': etudiant.nom},
                             'semestres': semestres_de(etudiant)})
        try:
            pdf_bytes, bilan = generer_releves_etudiants(etudiants, request.user)
        except ValueError as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        from django.http import HttpResponse
        resp = HttpResponse(pdf_bytes, content_type='application/pdf')
        nom = (f'releves_{etudiants[0].matricule}.pdf' if len(etudiants) == 1
               else f'releves_{len(etudiants)}_etudiants.pdf')
        resp['Content-Disposition'] = entete_piece_jointe(nom, inline=True)
        resp['X-Generated'] = str(sum(b['releves'] for b in bilan))
        resp['X-Total'] = str(sum(len(b['semestres']) for b in bilan))
        resp['X-Etudiants-Sans-Releve'] = ','.join(b['matricule'] for b in bilan if not b['releves'])
        resp['Access-Control-Expose-Headers'] = 'X-Generated, X-Total, X-Etudiants-Sans-Releve'
        return resp

    @action(detail=True, methods=['post'], url_path='regenerer')
    def regenerer(self, request, pk=None):
        """Vide le PDF cache d'un document — forcera la regeneration au prochain telechargement."""
        doc = self.get_object()
        _check_doc_module(request.user, doc.type_document, action='modifier')
        if doc.fichier_pdf:
            try:
                doc.fichier_pdf.delete(save=False)
            except Exception:
                pass
            doc.fichier_pdf = None
            doc.save(update_fields=['fichier_pdf'])
        return Response({'detail': 'PDF cache supprime, sera regenere au prochain telechargement.'})

    @action(detail=True, methods=['get'])
    def telecharger(self, request, pk=None):
        """Télécharge le PDF du document. 1re délivrance = original ; impressions
        suivantes = DUPLICATA (filigrane). Authentification requise."""
        from django.http import HttpResponse
        from django.utils import timezone
        doc = self.get_object()
        # Gate par type de doc : un user qui ne peut pas voir les diplomes
        # ne peut pas non plus telecharger le PDF d'un diplome.
        _check_doc_module(request.user, doc.type_document, action='voir')

        # is_duplicata décidé À LA DÉLIVRANCE : original tant que jamais délivré.
        is_duplicata = doc.premiere_generation is not None

        if not doc.fichier_pdf:
            from .services import _render_pdf, _get_qr_base64, _get_institution, _TEMPLATE_MAP, _CONTEXT_BUILDER, enregistrer_pdf_document
            from django.conf import settings
            etudiant = doc.etudiant
            institution = _get_institution(etudiant)
            base_url = getattr(settings, 'DOCUMENTS_BASE_URL', '').rstrip('/')
            template = _TEMPLATE_MAP.get(doc.type_document)
            builder = _CONTEXT_BUILDER.get(doc.type_document)
            if not template or not builder:
                return Response({'detail': f'Pas de template pour {doc.type_document}'}, status=status.HTTP_400_BAD_REQUEST)
            specific_ctx = builder(doc, etudiant, institution, {
                'annee_universitaire': doc.annee_universitaire,
                'semestre': doc.semestre_id,
            })
            qr_image = _get_qr_base64(
                str(doc.token_verification), base_url,
                etudiant=etudiant, doc=doc,
                filiere_nom=specific_ctx.get('filiere_nom', ''),
                niveau_label=specific_ctx.get('niveau_label', ''),
            )
            context = {
                'doc': doc,
                'etudiant': etudiant,
                'institution': institution or type('FakeInstitution', (), {
                    'nom': 'Université', 'nom_fr': '', 'nom_ar': '',
                    'logo': None, 'directeur_nom_fr': '', 'directeur_titre_fr': '', 'directeur_signature': None,
                })(),
                'today': timezone.now().strftime('%d/%m/%Y'),
                'verify_url': f'{base_url}/verifier/{doc.token_verification}',
                'qr_image': qr_image,
                'is_duplicata': is_duplicata,
                **specific_ctx,
            }
            try:
                pdf_bytes = _render_pdf(template, context)
            except Exception as e:
                return Response({'detail': f'Erreur génération PDF : {e}'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
            enregistrer_pdf_document(doc, pdf_bytes)

        # Lire les octets EN MÉMOIRE avant tout vidage de cache.
        with doc.fichier_pdf.open('rb') as f:
            pdf_bytes = f.read()
        filename = f'{doc.numero_serie}.pdf'

        # 1re délivrance : marquer le document comme délivré ET vider le cache, afin
        # que la PROCHAINE impression régénère un DUPLICATA (sinon le re-téléchargement
        # resservirait l'original mis en cache). Les duplicata restent en cache.
        if doc.premiere_generation is None:
            doc.premiere_generation = timezone.now()
            doc.fichier_pdf.delete(save=False)
            doc.fichier_pdf = None
            doc.save(update_fields=['premiere_generation', 'fichier_pdf'])

        response = HttpResponse(pdf_bytes, content_type='application/pdf')
        response['Content-Disposition'] = entete_piece_jointe(filename)
        # Anti-cache navigateur : sans ça, un re-téléchargement à la même URL peut
        # resservir l'ancien PDF gardé par le navigateur ("je vois pas de changement").
        response['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        response['Pragma'] = 'no-cache'
        return response

    @action(detail=False, methods=['get'], url_path='verifier/(?P<token>[^/.]+)',
            permission_classes=[AllowAny], throttle_classes=[VerificationThrottle])
    def verifier(self, request, token=None):
        """Vérifie un document par son token (public — AllowAny, rate-limité).

        TOUS les documents (attestation d'inscription, relevé, attestation de
        diplôme) → DocumentVerificationSerializer : identité + PHOTO + filière /
        mention / année, à COMPARER avec le papier. Aucune donnée technique (ni
        hash, ni chemin de fichier, ni auteur)."""
        from django.core.exceptions import ValidationError as DjangoValidationError
        try:
            doc = (DocumentOfficiel.objects
                   .select_related('etudiant', 'etudiant__filiere', 'institution')
                   .get(token_verification=token))
        except (DocumentOfficiel.DoesNotExist, ValueError, DjangoValidationError):
            # Polymorphe : une attestation enseignant vérifiable porte le même token.
            from .models import AttestationTravail
            try:
                att = (AttestationTravail.objects
                       .select_related('prof')
                       .get(token_verification=token))
            except (AttestationTravail.DoesNotExist, ValueError, DjangoValidationError):
                return Response({'detail': 'Document introuvable.'}, status=status.HTTP_404_NOT_FOUND)
            return Response(_serialize_attestation_travail(att))

        data = DocumentVerificationSerializer(doc, context={'request': request}).data
        return Response(data)

    @action(detail=False, methods=['get'], url_path='verifier/(?P<token>[^/.]+)/photo',
            permission_classes=[AllowAny], throttle_classes=[VerificationThrottle])
    def verifier_photo(self, request, token=None):
        """Photo de l'etudiant pour la page de verification PUBLIQUE (token-gated).

        Contourne l'auth_request /media/ (qui exige IsAuthenticated → la page de
        verification n'etant pas connectee, la photo /media/ renvoyait 401) SANS
        exposer les photos a l'enumeration : elle n'est servie que pour un token de
        verification VALIDE (UUID aleatoire). Meme exposition que l'identite deja
        affichee sur la page de verification."""
        import mimetypes
        from django.http import FileResponse
        from django.core.exceptions import ValidationError as DjangoValidationError
        try:
            doc = (DocumentOfficiel.objects
                   .select_related('etudiant')
                   .get(token_verification=token))
        except (DocumentOfficiel.DoesNotExist, ValueError, DjangoValidationError):
            return Response({'detail': 'Document introuvable.'}, status=status.HTTP_404_NOT_FOUND)

        photo = getattr(doc.etudiant, 'photo', None)
        if not photo:
            return Response({'detail': 'Aucune photo.'}, status=status.HTTP_404_NOT_FOUND)
        try:
            fh = photo.open('rb')
        except (FileNotFoundError, OSError):
            return Response({'detail': 'Photo indisponible.'}, status=status.HTTP_404_NOT_FOUND)

        ctype = mimetypes.guess_type(photo.name)[0] or 'application/octet-stream'
        resp = FileResponse(fh, content_type=ctype)
        # Semi-prive (accessible via token) : cache court cote client, jamais partage.
        resp['Cache-Control'] = 'private, max-age=300'
        resp['X-Content-Type-Options'] = 'nosniff'
        return resp



class RegistreDiplomeViewSet(InstitutionScopedMixin, viewsets.ReadOnlyModelViewSet):
    # Ordre de MÉRITE : moyenne décroissante (puis nom pour départager). S'applique
    # à la liste paginée ET à l'export. Filtrer par filière + année pour un
    # classement par promotion.
    queryset = (RegistreDiplome.objects
                .select_related('etudiant', 'filiere')
                .order_by('-moyenne_generale', 'etudiant__nom'))
    serializer_class   = RegistreDiplomeSerializer
    permission_classes = [RBACPermission]
    required_module    = 'doc_registre'
    filter_backends    = [DjangoFilterBackend, SearchFilter]
    filterset_fields   = ['filiere', 'annee_universitaire']
    search_fields      = ['numero_diplome', 'etudiant__nom', 'etudiant__matricule']

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """Export Excel du registre des diplômes."""
        import io
        try:
            import openpyxl
        except ImportError:
            return Response({'detail': 'openpyxl non installé.'}, status=status.HTTP_501_NOT_IMPLEMENTED)
        from django.http import HttpResponse

        qs = self.filter_queryset(self.get_queryset())
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Registre diplômes'
        headers = ['Rang', 'N° Diplôme', 'Matricule', 'Nom', 'Filière', 'Moyenne', 'Mention', 'Année', 'Date']
        ws.append(headers)
        for rang, r in enumerate(qs, start=1):
            ws.append([
                rang, r.numero_diplome, r.etudiant.matricule, r.etudiant.nom,
                r.filiere.intitule_fr, float(r.moyenne_generale),
                r.mention, r.annee_universitaire, str(r.date_delivrance),
            ])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        resp = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = entete_piece_jointe('registre_diplomes.xlsx')
        return resp

    @action(detail=False, methods=['get'], url_path='export-ministere')
    def export_ministere(self, request):
        """Export Excel des diplômés Licence L3 selon la maquette demandée par le
        Ministère (MESRS) : Establishment | Filière L3 | NNI | Numéro d'inscription
        | Prénom Nom | Moyenne Cumulative | Rang.

        - Source = registre des diplômes (fin de cycle) scopé à la Licence L3
          (filiere.niveau_fin == 3) ; respecte les filtres filière + année.
        - NNI = etudiant.cni (même source que l'attestation de diplôme).
        - Moyenne Cumulative = moyenne_generale (moyenne des 6 semestres du cycle).
        - Rang = rang de mérite PAR FILIÈRE (promotion), moyenne décroissante.
        """
        import io
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            return Response({'detail': 'openpyxl non installé.'}, status=status.HTTP_501_NOT_IMPLEMENTED)
        from django.http import HttpResponse

        qs = (self.filter_queryset(self.get_queryset())
              .filter(filiere__niveau_fin=3)   # Licence L3 (fin de cycle)
              .select_related('etudiant', 'filiere', 'institution'))

        # Rang de mérite PAR FILIÈRE : tri (filière, moyenne décroissante, nom).
        lignes = sorted(qs, key=lambda x: (
            x.filiere.intitule_fr or x.filiere.code or '',
            -float(x.moyenne_generale or 0),
            (x.etudiant.nom_fr or x.etudiant.nom or ''),
        ))
        rang_par_filiere = {}

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Diplômés Licence L3'
        headers = ['Establishment', 'Filière L3', 'NNI', "Numéro d'inscription",
                   'Prénom Nom', 'Moyenne Cumulative', 'Rang']
        ws.append(headers)

        PRIMARY = '006633'
        thin = Side(border_style='thin', color='CCCCCC')
        border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor=PRIMARY)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = border_all

        for r in lignes:
            etu  = r.etudiant
            inst = r.institution
            etab = (getattr(inst, 'nom_fr', '') or getattr(inst, 'nom', '') or '')
            fil  = r.filiere.intitule_fr or r.filiere.code or ''
            prenom_nom = f"{etu.prenom_fr or ''} {etu.nom_fr or etu.nom or ''}".strip()
            rang_par_filiere[r.filiere_id] = rang_par_filiere.get(r.filiere_id, 0) + 1
            ws.append([
                etab, fil, etu.cni or '', etu.matricule,
                prenom_nom, float(r.moyenne_generale), rang_par_filiere[r.filiere_id],
            ])
            for c in ws[ws.max_row]:
                c.border = border_all
            ws.cell(ws.max_row, 6).alignment = Alignment(horizontal='center')
            ws.cell(ws.max_row, 7).alignment = Alignment(horizontal='center')

        for col, w in zip('ABCDEFG', [30, 28, 16, 18, 32, 18, 8]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = 'A2'

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        resp = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = entete_piece_jointe('diplomes_licence_L3_ministere.xlsx')
        return resp
