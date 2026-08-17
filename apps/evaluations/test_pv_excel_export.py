"""
Non-régression de l'export Excel du PV de délibération.

Garde-fou avant/après l'extraction de la méthode PVDeliberationViewSet.excel
(~500 lignes) vers le service services/pv_export.build_pv_excel. L'endpoint doit
continuer à renvoyer un classeur xlsx valide à structure attendue.

Base SQLite en mémoire (siga.settings.test) — aucun contact gesafped26/siga.
"""
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db import connection
from rest_framework.test import APITestCase
from openpyxl import load_workbook

from apps.prof.models import ProfTypeHistory
from tests.factories.deliberation import PVDeliberationSemestrielFactory

User = get_user_model()


class PVExcelExportTest(APITestCase):
    @classmethod
    def setUpClass(cls):
        # create_model plante la transaction PG si la table existe deja (creee par
        # la fixture session de tests/conftest.py) : verifier l'existence AVANT
        # d'entrer dans le schema_editor est vendor-neutre (sqlite ET postgresql).
        if ProfTypeHistory._meta.db_table not in connection.introspection.table_names():
            with connection.schema_editor() as se:
                se.create_model(ProfTypeHistory)
        super().setUpClass()

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_user(
            username='admin_pv', email='admin_pv@test.local',
            password='x', role='admin',
        )
        # InstitutionFactory (get_or_create acronyme + est_principale=True) →
        # le PV et ses sous-objets partagent l'unique institution principale,
        # donc InstitutionScopedMixin.get_object() le trouve bien.
        cls.pv = PVDeliberationSemestrielFactory()

    def test_excel_export_renvoie_classeur_valide(self):
        self.client.force_authenticate(self.admin)
        resp = self.client.get(f'/api/v1/evaluations/pvs/{self.pv.id}/excel/')

        self.assertEqual(resp.status_code, 200, getattr(resp, 'content', b'')[:300])
        self.assertIn('spreadsheet', resp['Content-Type'])

        wb = load_workbook(BytesIO(resp.content))
        self.assertGreaterEqual(len(wb.sheetnames), 1)
        # PV semestriel sur session normale → la feuille "Matrice rattrapages" existe.
        self.assertIn('Matrice rattrapages', wb.sheetnames)
