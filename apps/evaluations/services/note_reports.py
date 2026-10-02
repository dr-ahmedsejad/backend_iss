"""Rapports/fiches PDF du domaine notes (émargement, collecte, levée
d'anonymat) + gestion des anonymats — extraits des ViewSets reports.
"""
import logging
from datetime import date

from django.http import HttpResponse
from django.template.loader import get_template
from rest_framework import status
from rest_framework.response import Response
from core.telechargement import entete_piece_jointe

logger = logging.getLogger('siga')

from apps.evaluations.models import (
    SessionEvaluation, Note, ResultatElement, ResultatSemestre,
    ObligationRattrapage, ResultatModule,
)
from ..views_helpers import (
    _build_institution_context, _render_pdf,
    _enrichir_lignes_anonymat, _resolve_em,
)


def _emargement_etudiants(filiere_id, niveau, semestre, annee_univ):
    """Liste PLATE des étudiants inscrits (filière + niveau + semestre + année),
    triée par matricule. Source UNIQUE partagée par l'émargement PDF et Excel."""
    from apps.inscriptions.models import InscriptionPedagogique
    insc_peds = InscriptionPedagogique.objects.filter(
        semestre__code_semestre=semestre,
        inscription_admin__filiere_id=filiere_id,
        inscription_admin__niveau=niveau,
        inscription_admin__annee_univ_id=annee_univ,
    ).select_related('inscription_admin__etudiant').order_by(
        'inscription_admin__etudiant__matricule',
    )
    return [ip.inscription_admin.etudiant for ip in insc_peds]


def build_emargement_pdf(request):
    filiere_id  = request.query_params.get('filiere')
    niveau      = request.query_params.get('niveau')
    semestre    = request.query_params.get('semestre')
    annee_univ  = request.query_params.get('annee_univ')

    if not all([filiere_id, niveau, semestre, annee_univ]):
        return Response(
            {'detail': 'Paramètres requis : filiere, niveau, semestre, annee_univ.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    # Liste PLATE par matricule (source unique partagée avec l'émargement Excel),
    # sans découpage par département — aligné sur les fiches de collecte.
    etudiants = _emargement_etudiants(filiere_id, niveau, semestre, annee_univ)
    groupes = [{'nom': '', 'etudiants': etudiants}]

    from apps.scolarite.models import Filiere
    from apps.parametres.models import Year
    try:
        filiere = Filiere.objects.get(pk=filiere_id)
    except Filiere.DoesNotExist:
        filiere = None
    try:
        annee = Year.objects.get(pk=annee_univ)
    except Year.DoesNotExist:
        annee = None

    institution, institution_logo_url = _build_institution_context()

    # En-tête commun aux fiches d'absence (suivi_en_tete.html) : on injecte
    # le contexte institution standard + l'année universitaire.
    from core.pdf_utils import get_institution_context
    context = {
        **get_institution_context(),
        'annee_universitaire':  annee.annee if annee else '',
        'groupes':              groupes,
        'etudiants':            etudiants,
        'filiere':              filiere,
        'niveau':               niveau,
        'semestre':             semestre,
        'annee':                annee,
        'institution':          institution,
        'institution_logo_url': institution_logo_url,
        'date_impression':      date.today().strftime('%d/%m/%Y'),
        'nb_etudiants':         len(etudiants),
    }
    fn = f'emargement_{(filiere.code if filiere else filiere_id)}_N{niveau}_{semestre}.pdf'
    return _render_pdf('fiche_emargement.html', context, fn)


def build_emargement_excel(request):
    """Version EXCEL de la fiche d'émargement — MÊME liste que le PDF (plate,
    triée par matricule). Colonnes : Matricule | Nom et Prénom | Signature (vide)."""
    import io
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return Response({'detail': 'openpyxl non installé.'}, status=status.HTTP_501_NOT_IMPLEMENTED)
    from django.http import HttpResponse

    filiere_id  = request.query_params.get('filiere')
    niveau      = request.query_params.get('niveau')
    semestre    = request.query_params.get('semestre')
    annee_univ  = request.query_params.get('annee_univ')
    if not all([filiere_id, niveau, semestre, annee_univ]):
        return Response(
            {'detail': 'Paramètres requis : filiere, niveau, semestre, annee_univ.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    etudiants = _emargement_etudiants(filiere_id, niveau, semestre, annee_univ)

    from apps.scolarite.models import Filiere
    from apps.parametres.models import Year
    filiere = Filiere.objects.filter(pk=filiere_id).first()
    annee = Year.objects.filter(pk=annee_univ).first()

    PRIMARY = '006633'
    thin = Side(border_style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hfont = Font(bold=True, color='FFFFFF')
    hfill = PatternFill('solid', fgColor=PRIMARY)
    center = Alignment(horizontal='center', vertical='center')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Emargement'
    ws['A1'] = "Fiche d'Émargement"
    ws['A1'].font = Font(bold=True, size=12, color=PRIMARY)
    ws['A2'] = f"Filière : {filiere.intitule_fr if filiere else ''}"
    ws['A3'] = f"Niveau : L{niveau}    Semestre : {semestre}    Année : {annee.annee if annee else ''}"

    hrow = 5
    for col, h in enumerate(['Matricule', 'Nom et Prénom', 'Signature'], 1):
        c = ws.cell(hrow, col, h)
        c.font = hfont
        c.fill = hfill
        c.alignment = center
        c.border = border

    r = hrow + 1
    for etu in etudiants:
        ws.cell(r, 1, etu.matricule).font = Font(name='Consolas', size=10)
        ws.cell(r, 2, etu.nom_fr or etu.nom or '')
        ws.cell(r, 3, None)
        for col in range(1, 4):
            ws.cell(r, col).border = border
        r += 1

    for idx, w in enumerate([16, 34, 30], 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = ws.cell(hrow + 1, 1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    fn = f'emargement_{(filiere.code if filiere else filiere_id)}_N{niveau}_{semestre}.xlsx'
    resp['Content-Disposition'] = entete_piece_jointe(fn)
    return resp


def build_collecte_pdf(request, type_session_attendu):
    from apps.inscriptions.models import InscriptionElement

    em_id      = request.query_params.get('em')
    session_id = request.query_params.get('session')
    type_note  = request.query_params.get('type_note', 'EXAM').upper()
    anonymat   = request.query_params.get('anonymat', '0') == '1'

    if not em_id or not session_id:
        return Response(
            {'detail': 'Paramètres requis : em, session.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        session = SessionEvaluation.objects.select_related('annee_univ').get(pk=session_id)
    except SessionEvaluation.DoesNotExist:
        return Response({'detail': 'Session introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    if session.type_session != type_session_attendu:
        return Response(
            {'detail': f'Cette fiche nécessite une session de type "{type_session_attendu}".'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    em = _resolve_em(em_id)
    if not em:
        return Response({'detail': 'EM introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    if type_note == 'TP' and not getattr(em, 'has_tp', False):
        return Response(
            {'detail': 'Cet EM n\'a pas de TP (has_tp=False).'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    # Liste des étudiants — ALIGNÉE avec la feuille de saisie (build_feuille) :
    # MÊME ensemble d'InscriptionElement et MÊME ordre (matricule). Source unique
    # partagée avec l'export Excel (cf. _resolve_collecte_lignes).
    lignes_etudiants = _resolve_collecte_lignes(em, session, type_session_attendu)

    # Anonymat
    anonymat_map = {}
    if anonymat:
        from apps.evaluations.models import AnonymatSession
        anons = AnonymatSession.objects.filter(session=session)
        if not anons.exists():
            return Response(
                {'detail': 'Générer d\'abord les anonymats de cette session.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        anonymat_map = {a.inscription_admin_id: a.numero_anonymat for a in anons}
        lignes_etudiants = _enrichir_lignes_anonymat(lignes_etudiants, anonymat_map)

    TYPE_NOTE_LABELS = {'CC': 'Contrôle Continu', 'TP': 'Travaux Pratiques', 'EXAM': 'Examen Final'}
    institution, institution_logo_url = _build_institution_context()

    is_rattrapage = (type_session_attendu == 'rattrapage')
    template = 'fiche_collecte_rattrapage.html' if is_rattrapage else 'fiche_collecte_notes.html'

    # Libellé du type de note : pour un examen de session rattrapage,
    # « Examen de Rattrapage » au lieu de « Examen Final ».
    type_note_label = TYPE_NOTE_LABELS.get(type_note, type_note)
    if type_note == 'EXAM' and is_rattrapage:
        type_note_label = 'Examen de Rattrapage'

    # Alignement STRICT avec la feuille de saisie : liste PLATE par matricule, sans
    # découpage par groupe ; filière dérivée de l'inscription administrative
    # (cf. _collecte_filiere_nom).
    groupes = [{'nom': '', 'filiere': _collecte_filiere_nom(lignes_etudiants), 'lignes': lignes_etudiants}]

    from core.pdf_utils import get_institution_context
    context = {
        **get_institution_context(),
        'annee_universitaire':  session.annee_univ.annee if session.annee_univ_id else '',
        'em':                   em,
        'session':              session,
        'type_note':            type_note,
        'type_note_label':      type_note_label,
        'groupes':              groupes,
        'lignes':               lignes_etudiants,
        'anonymat':             anonymat,
        'anonymat_map':         anonymat_map,
        'institution':          institution,
        'institution_logo_url': institution_logo_url,
        'date_impression':      date.today().strftime('%d/%m/%Y'),
        'nb_lignes':            len(lignes_etudiants),
        'is_rattrapage':        is_rattrapage,
    }
    suffix = 'rattrapage' if is_rattrapage else 'normale'
    fn = f'collecte_{em.code_em}_{type_note}_{suffix}.pdf'
    return _render_pdf(template, context, fn)


def build_collecte_pdf_tous(request, type_session_attendu):
    from apps.em.models import EM
    from apps.inscriptions.models import InscriptionElement
    from apps.evaluations.models import ObligationRattrapage, AnonymatSession

    session_id  = request.query_params.get('session')
    filiere_id  = request.query_params.get('filiere')
    semestre_id = request.query_params.get('semestre')
    type_note   = request.query_params.get('type_note', 'EXAM').upper()
    anonymat    = request.query_params.get('anonymat', '0') == '1'

    if not all([session_id, filiere_id, semestre_id]):
        return Response(
            {'detail': 'Paramètres requis : session, filiere, semestre.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        session = SessionEvaluation.objects.select_related('annee_univ').get(pk=session_id)
    except SessionEvaluation.DoesNotExist:
        return Response({'detail': 'Session introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    if session.type_session != type_session_attendu:
        return Response(
            {'detail': f'Cette fiche nécessite une session de type "{type_session_attendu}".'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    ems = EM.objects.filter(
        semestre_id=semestre_id,
        inscriptions_elements__inscription_ped__inscription_admin__filiere_id=filiere_id,
        inscriptions_elements__inscription_ped__inscription_admin__annee_univ=session.annee_univ,
    ).distinct().order_by('code_em')

    if type_note == 'TP':
        ems = ems.filter(has_tp=True)

    # Anonymat map global
    anonymat_map = {}
    if anonymat:
        anons = AnonymatSession.objects.filter(session=session)
        if not anons.exists():
            return Response(
                {'detail': 'Générer d\'abord les anonymats de cette session.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        anonymat_map = {a.inscription_admin_id: a.numero_anonymat for a in anons}

    from apps.scolarite.models import Filiere
    _fil = Filiere.objects.filter(pk=filiere_id).first()
    filiere_nom = _fil.intitule_fr if _fil else ''

    ems_data = []
    for em in ems:
        # Même liste que le PDF et que la feuille de saisie (source unique partagée).
        lignes = _resolve_collecte_lignes(em, session, type_session_attendu)
        if anonymat:
            lignes = _enrichir_lignes_anonymat(lignes, anonymat_map)

        ems_data.append({
            'em':          em,
            'lignes':      lignes,
            'nb_lignes':   len(lignes),
            'anonymat_map': anonymat_map,
            # Alignement avec la feuille de saisie : liste PLATE par matricule (pas
            # de découpage par groupe), pour que la fiche corresponde à l'écran.
            'groupes':     [{'nom': '', 'filiere': filiere_nom, 'lignes': lignes}],
        })

    TYPE_NOTE_LABELS = {'CC': 'Contrôle Continu', 'TP': 'Travaux Pratiques', 'EXAM': 'Examen Final'}
    institution, institution_logo_url = _build_institution_context()
    suffix = 'rattrapage' if type_session_attendu == 'rattrapage' else 'normale'

    # Examen de session rattrapage → « Examen de Rattrapage ».
    type_note_label = TYPE_NOTE_LABELS.get(type_note, type_note)
    if type_note == 'EXAM' and type_session_attendu == 'rattrapage':
        type_note_label = 'Examen de Rattrapage'

    from core.pdf_utils import get_institution_context
    context = {
        **get_institution_context(),
        'annee_universitaire':  session.annee_univ.annee if session.annee_univ_id else '',
        'ems_data':             ems_data,
        'session':              session,
        'type_note':            type_note,
        'type_note_label':      type_note_label,
        'anonymat':             anonymat,
        'institution':          institution,
        'institution_logo_url': institution_logo_url,
        'date_impression':      date.today().strftime('%d/%m/%Y'),
    }
    fn = f'collecte_tous_{type_note}_{suffix}.pdf'
    return _render_pdf('fiche_collecte_tous.html', context, fn)


# ── Résolution PARTAGÉE des étudiants d'une fiche de collecte ─────────────────
# Source UNIQUE utilisée par les exports PDF ET Excel (individuel + tous), pour
# que PDF et Excel affichent EXACTEMENT la même liste — alignée sur la feuille de
# saisie (build_feuille) : même ensemble d'InscriptionElement + tri matricule.

_TYPE_NOTE_LABELS = {'CC': 'Contrôle Continu', 'TP': 'Travaux Pratiques', 'EXAM': 'Examen Final'}


def _type_note_label(type_note, is_rattrapage) -> str:
    if type_note == 'EXAM' and is_rattrapage:
        return 'Examen de Rattrapage'
    return _TYPE_NOTE_LABELS.get(type_note, type_note)


def _resolve_collecte_lignes(em, session, type_session_attendu) -> list:
    """[{'etudiant','inscription_admin','type_obligation'}] triés par matricule.
    Rattrapage → eligible_rattrapage_ie_ids (obligations gardées ∪ dettes non
    validées, source unique partagée avec la saisie et l'import) ; normale →
    inscrits à l'EM pour l'année de la session."""
    from apps.inscriptions.models import InscriptionElement
    sel = (
        'inscription_ped__inscription_admin__etudiant',
        'inscription_ped__inscription_admin__filiere',
    )
    order = 'inscription_ped__inscription_admin__etudiant__matricule'
    if type_session_attendu == 'rattrapage':
        from apps.evaluations.services.note_lecture import eligible_rattrapage_ie_ids
        from apps.evaluations.models import ObligationRattrapage
        ie_ids = list(eligible_rattrapage_ie_ids(em.id, session))
        type_oblig_map = dict(ObligationRattrapage.objects.filter(
            inscription_element_id__in=ie_ids,
            ligne__pv__session__annee_univ=session.annee_univ,
        ).values_list('inscription_element_id', 'type_obligation'))
        ies = InscriptionElement.objects.filter(id__in=ie_ids).select_related(*sel).order_by(order)
        return [{
            'etudiant':          ie.inscription_ped.inscription_admin.etudiant,
            'inscription_admin': ie.inscription_ped.inscription_admin,
            'type_obligation':   type_oblig_map.get(ie.id),
        } for ie in ies]
    ies = InscriptionElement.objects.filter(
        em=em,
        inscription_ped__inscription_admin__annee_univ=session.annee_univ,
    ).select_related(*sel).order_by(order)
    return [{
        'etudiant':          ie.inscription_ped.inscription_admin.etudiant,
        'inscription_admin': ie.inscription_ped.inscription_admin,
        'type_obligation':   None,
    } for ie in ies]


def _collecte_filiere_nom(lignes_etudiants) -> str:
    """Filière(s) affichée(s) — dérivée de l'inscription administrative (vraie
    filière de l'année, PAS le département du tronc commun), distinctes jointes si
    l'EM est mutualisé entre filières."""
    fils = []
    for lg in lignes_etudiants:
        ia = lg.get('inscription_admin')
        fn = (ia.filiere.intitule_fr if ia and getattr(ia, 'filiere_id', None) else '')
        if fn and fn not in fils:
            fils.append(fn)
    return ' / '.join(fils)


def _collecte_excel_response(ems_data, session, type_note_label, anonymat, filename):
    """Classeur Excel de collecte : une feuille par EM, MÊMES étudiants que le PDF.
    Colonnes : Matricule/Nom (ou N° Anonymat) + Note /20 (vide, à remplir) + Observations."""
    import io
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return Response({'detail': 'openpyxl non installé.'}, status=status.HTTP_501_NOT_IMPLEMENTED)
    from django.http import HttpResponse

    PRIMARY = '006633'
    thin = Side(border_style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hfont = Font(bold=True, color='FFFFFF')
    hfill = PatternFill('solid', fgColor=PRIMARY)
    center = Alignment(horizontal='center', vertical='center')
    title_font = Font(bold=True, size=12, color=PRIMARY)
    annee = session.annee_univ.annee if session.annee_univ_id else ''

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for data in ems_data:
        em = data['em']
        lignes = data['lignes']
        base = (em.code_em or 'EM')[:28]
        name = base
        k = 1
        while name in wb.sheetnames:
            k += 1
            name = f'{base}_{k}'
        ws = wb.create_sheet(name)

        ws['A1'] = f'Fiche de collecte — {type_note_label}'
        ws['A1'].font = title_font
        ws['A2'] = f'EM : {em.code_em} — {em.intitule}'
        ws['A3'] = f"Filière : {data.get('filiere_nom', '')}"
        sem_code = em.semestre.code_semestre if getattr(em, 'semestre_id', None) else ''
        ws['A4'] = f'Semestre : {sem_code}    Session : {session.code or ""}    Année : {annee}'

        headers = (['N° Anonymat', 'Note /20', 'Observations'] if anonymat
                   else ['Matricule', 'Nom et Prénom', 'Note /20', 'Observations'])
        hrow = 6
        for col, h in enumerate(headers, 1):
            c = ws.cell(hrow, col, h)
            c.font = hfont
            c.fill = hfill
            c.alignment = center
            c.border = border

        r = hrow + 1
        for lg in lignes:
            etu = lg['etudiant']
            vals = ([lg.get('numero_anonymat', ''), None, None] if anonymat
                    else [etu.matricule, etu.nom_fr or etu.nom or '', None, None])
            for col, v in enumerate(vals, 1):
                cell = ws.cell(r, col, v)
                cell.border = border
                if col == 1 and not anonymat:
                    cell.font = Font(name='Consolas', size=10)
            r += 1

        widths = ([16, 12, 34] if anonymat else [16, 34, 12, 34])
        for idx, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        ws.freeze_panes = ws.cell(hrow + 1, 1)

    if not wb.sheetnames:
        wb.create_sheet('Vide')

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resp['Content-Disposition'] = entete_piece_jointe(filename)
    return resp


def build_collecte_excel(request, type_session_attendu):
    """Version EXCEL de la fiche de collecte individuelle (mêmes étudiants que le PDF)."""
    em_id      = request.query_params.get('em')
    session_id = request.query_params.get('session')
    type_note  = request.query_params.get('type_note', 'EXAM').upper()
    anonymat   = request.query_params.get('anonymat', '0') == '1'

    if not em_id or not session_id:
        return Response({'detail': 'Paramètres requis : em, session.'}, status=status.HTTP_400_BAD_REQUEST)
    try:
        session = SessionEvaluation.objects.select_related('annee_univ').get(pk=session_id)
    except SessionEvaluation.DoesNotExist:
        return Response({'detail': 'Session introuvable.'}, status=status.HTTP_404_NOT_FOUND)
    if session.type_session != type_session_attendu:
        return Response({'detail': f'Cette fiche nécessite une session de type "{type_session_attendu}".'},
                        status=status.HTTP_400_BAD_REQUEST)
    em = _resolve_em(em_id)
    if not em:
        return Response({'detail': 'EM introuvable.'}, status=status.HTTP_404_NOT_FOUND)
    if type_note == 'TP' and not getattr(em, 'has_tp', False):
        return Response({'detail': "Cet EM n'a pas de TP (has_tp=False)."}, status=status.HTTP_400_BAD_REQUEST)

    lignes = _resolve_collecte_lignes(em, session, type_session_attendu)
    if anonymat:
        from apps.evaluations.models import AnonymatSession
        anons = AnonymatSession.objects.filter(session=session)
        if not anons.exists():
            return Response({'detail': "Générer d'abord les anonymats de cette session."},
                            status=status.HTTP_400_BAD_REQUEST)
        lignes = _enrichir_lignes_anonymat(
            lignes, {a.inscription_admin_id: a.numero_anonymat for a in anons})

    is_ratt = (type_session_attendu == 'rattrapage')
    ems_data = [{'em': em, 'lignes': lignes, 'filiere_nom': _collecte_filiere_nom(lignes)}]
    suffix = 'rattrapage' if is_ratt else 'normale'
    return _collecte_excel_response(
        ems_data, session, _type_note_label(type_note, is_ratt), anonymat,
        f'collecte_{em.code_em}_{type_note}_{suffix}.xlsx')


def build_collecte_excel_tous(request, type_session_attendu):
    """Version EXCEL de la fiche « tous les EM » (une feuille par EM)."""
    from apps.em.models import EM
    from apps.evaluations.models import AnonymatSession
    from apps.scolarite.models import Filiere

    session_id  = request.query_params.get('session')
    filiere_id  = request.query_params.get('filiere')
    semestre_id = request.query_params.get('semestre')
    type_note   = request.query_params.get('type_note', 'EXAM').upper()
    anonymat    = request.query_params.get('anonymat', '0') == '1'

    if not all([session_id, filiere_id, semestre_id]):
        return Response({'detail': 'Paramètres requis : session, filiere, semestre.'},
                        status=status.HTTP_400_BAD_REQUEST)
    try:
        session = SessionEvaluation.objects.select_related('annee_univ').get(pk=session_id)
    except SessionEvaluation.DoesNotExist:
        return Response({'detail': 'Session introuvable.'}, status=status.HTTP_404_NOT_FOUND)
    if session.type_session != type_session_attendu:
        return Response({'detail': f'Cette fiche nécessite une session de type "{type_session_attendu}".'},
                        status=status.HTTP_400_BAD_REQUEST)

    ems = EM.objects.filter(
        semestre_id=semestre_id,
        inscriptions_elements__inscription_ped__inscription_admin__filiere_id=filiere_id,
        inscriptions_elements__inscription_ped__inscription_admin__annee_univ=session.annee_univ,
    ).distinct().order_by('code_em')
    if type_note == 'TP':
        ems = ems.filter(has_tp=True)

    anonymat_map = {}
    if anonymat:
        anons = AnonymatSession.objects.filter(session=session)
        if not anons.exists():
            return Response({'detail': "Générer d'abord les anonymats de cette session."},
                            status=status.HTTP_400_BAD_REQUEST)
        anonymat_map = {a.inscription_admin_id: a.numero_anonymat for a in anons}

    _fil = Filiere.objects.filter(pk=filiere_id).first()
    filiere_nom = _fil.intitule_fr if _fil else ''

    ems_data = []
    for em in ems:
        lignes = _resolve_collecte_lignes(em, session, type_session_attendu)
        if anonymat:
            lignes = _enrichir_lignes_anonymat(lignes, anonymat_map)
        ems_data.append({'em': em, 'lignes': lignes, 'filiere_nom': filiere_nom})

    is_ratt = (type_session_attendu == 'rattrapage')
    suffix = 'rattrapage' if is_ratt else 'normale'
    return _collecte_excel_response(
        ems_data, session, _type_note_label(type_note, is_ratt), anonymat,
        f'collecte_tous_{type_note}_{suffix}.xlsx')


def list_anonymats(request):
    from apps.evaluations.models import AnonymatSession
    from apps.evaluations.serializers import AnonymatSessionSerializer
    session_id = request.query_params.get('session')
    if not session_id:
        return Response({'detail': 'Paramètre session requis.'}, status=status.HTTP_400_BAD_REQUEST)
    qs = AnonymatSession.objects.filter(
        session_id=session_id,
    ).select_related('inscription_admin__etudiant').order_by('numero_anonymat')
    serializer = AnonymatSessionSerializer(qs, many=True)
    return Response({'count': qs.count(), 'results': serializer.data})


def do_generer_anonymat(request):
    from apps.evaluations.services.anonymat import AnonymatService
    session_id = request.query_params.get('session')
    regenerer  = request.query_params.get('regenerer', '0') == '1'
    # Override du garde-fou "notes déjà saisies" : réservé à un admin.
    u = request.user
    force = (request.query_params.get('force', '0') == '1'
             and (getattr(u, 'is_superuser', False) or getattr(u, 'role', None) == 'admin'))
    if not session_id:
        return Response({'detail': 'Paramètre session requis.'}, status=status.HTTP_400_BAD_REQUEST)
    try:
        session = SessionEvaluation.objects.get(pk=session_id)
    except SessionEvaluation.DoesNotExist:
        return Response({'detail': 'Session introuvable.'}, status=status.HTTP_404_NOT_FOUND)
    try:
        nb = AnonymatService.generer(session, regenerer=regenerer,
                                     genere_par=request.user, force=force)
    except ValueError as e:
        return Response({'detail': str(e)}, status=status.HTTP_409_CONFLICT)

    # Audit : tracer qui (ré)génère les anonymats, pour quelle session.
    from core.audit_helpers import write_audit
    from core.models import ACTION_CREATE, ACTION_UPDATE
    write_audit(
        action=ACTION_UPDATE if regenerer else ACTION_CREATE,
        model_name='AnonymatSession',
        object_id=str(session.pk),
        changes={
            'session':       session.intitule or session.code or str(session.pk),
            'regeneration':  regenerer,
            'force':         force,
            'nb_generes':    nb,
        },
        label=f"{'Régénération' if regenerer else 'Génération'} anonymats — "
              f"session {session.intitule or session.code or session.pk} ({nb} numéros)",
    )
    return Response({'nb_generes': nb})


def build_levee_pdf(request):
    from apps.evaluations.models import AnonymatSession
    session_id = request.query_params.get('session')
    if not session_id:
        return Response({'detail': 'Paramètre session requis.'}, status=status.HTTP_400_BAD_REQUEST)
    try:
        session = SessionEvaluation.objects.select_related('annee_univ').get(pk=session_id)
    except SessionEvaluation.DoesNotExist:
        return Response({'detail': 'Session introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    anonymats = AnonymatSession.objects.filter(
        session=session,
    ).select_related('inscription_admin__etudiant').order_by('numero_anonymat')

    if not anonymats.exists():
        return Response(
            {'detail': 'Aucun anonymat généré pour cette session.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    institution, institution_logo_url = _build_institution_context()
    context = {
        'session':              session,
        'anonymats':            list(anonymats),
        'institution':          institution,
        'institution_logo_url': institution_logo_url,
        'date_impression':      date.today().strftime('%d/%m/%Y'),
    }
    fn = f'levee_anonymat_{session.code or session_id}.pdf'
    return _render_pdf('fiche_levee_anonymat.html', context, fn)

