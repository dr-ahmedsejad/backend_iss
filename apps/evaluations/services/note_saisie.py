"""Écritures du domaine notes (saisie anonyme / import xlsx / saisie en lot).

Extraites de NoteViewSet. Le calcul lourd reste dans NoteCalculService.
"""
from rest_framework import status
from rest_framework.response import Response

from apps.evaluations.models import Note, SessionEvaluation
from apps.evaluations.services.calcul_notes import NoteCalculService
from .note_access import peut_acceder_em


def do_saisir_anonymat(request):
    """
    POST {session, em, numero_anonymat, valeur, type_note}
    Résout le numéro d'anonymat → InscriptionElement et enregistre la note.
    """
    from decimal import Decimal, InvalidOperation
    from apps.evaluations.services.anonymat import AnonymatService
    from apps.evaluations.models import AnonymatSession as _AS
    from apps.inscriptions.models import InscriptionElement

    session_id      = request.data.get('session')
    em_id           = request.data.get('em')
    numero_anonymat = request.data.get('numero_anonymat')
    valeur_raw      = request.data.get('valeur')
    type_note       = str(request.data.get('type_note', 'EXAM')).upper()

    if not all([session_id, em_id, numero_anonymat is not None, valeur_raw is not None]):
        return Response({'detail': 'session, em, numero_anonymat, valeur requis.'},
                        status=status.HTTP_400_BAD_REQUEST)

    # Scope enseignant : un prof ne saisit que pour SES EMs (cf. saisir_bulk).
    # Les rôles all-EM (admin/scolarite/DE via eval_saisie) passent ce contrôle.
    if not peut_acceder_em(request.user, em_id):
        return Response({'detail': "Accès non autorisé : saisie sur un EM que vous n'enseignez pas."},
                        status=status.HTTP_403_FORBIDDEN)

    try:
        session = SessionEvaluation.objects.get(pk=session_id)
    except SessionEvaluation.DoesNotExist:
        return Response({'detail': 'Session introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    if session.est_close:
        return Response({'detail': 'Session clôturée — saisie impossible.'},
                        status=status.HTTP_400_BAD_REQUEST)

    try:
        valeur = Decimal(str(valeur_raw).replace(',', '.'))
    except InvalidOperation:
        return Response({'detail': 'Note invalide.'}, status=status.HTTP_400_BAD_REQUEST)

    if not (Decimal('0') <= valeur <= Decimal('20')):
        return Response({'detail': 'La note doit être entre 0 et 20.'},
                        status=status.HTTP_400_BAD_REQUEST)

    if type_note not in ('CC', 'TP', 'EXAM'):
        return Response({'detail': 'type_note doit être CC, TP ou EXAM.'},
                        status=status.HTTP_400_BAD_REQUEST)

    # Résoudre anonymat → InscriptionAdministrative
    try:
        ia = AnonymatService.resoudre(session, int(numero_anonymat))
    except _AS.DoesNotExist:
        return Response({'detail': f"Numéro d'anonymat {numero_anonymat} introuvable pour cette session."},
                        status=status.HTTP_404_NOT_FOUND)

    em = _resolve_em(em_id)
    if not em:
        return Response({'detail': 'EM introuvable.'}, status=status.HTTP_404_NOT_FOUND)

    ie = InscriptionElement.objects.filter(
        em=em,
        inscription_ped__inscription_admin=ia,
    ).first()
    if not ie:
        return Response({'detail': "Cet étudiant n'est pas inscrit à cet EM."},
                        status=status.HTTP_404_NOT_FOUND)

    note, created = Note.objects.update_or_create(
        inscription_element=ie,
        session=session,
        type_note=type_note,
        defaults={'valeur': valeur, 'saisie_par': request.user},
    )
    return Response({
        'id':                  note.id,
        'created':             created,
        'numero_anonymat':     int(numero_anonymat),
        'type_note':           type_note,
        'valeur':              float(valeur),
        'inscription_element': ie.id,
    }, status=status.HTTP_200_OK)


def _normalize_matricule(raw):
    """Normalise un matricule lu d'Excel (enlève le '.0' final si float, strip)."""
    if raw is None:
        return ''
    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw))
    s = str(raw).strip()
    if s.endswith('.0') and s[:-2].isdigit():
        s = s[:-2]
    return s


def _resoudre_col_map(headers, type_note, type_session):
    """type_note -> nom de colonne xlsx. Alias EXAM pour le rattrapage (Art. 18 :
    seul l'EXAM change en rattrapage)."""
    EXAM_ALIASES = ['note_exam', 'note_rat', 'note_rattrapage', 'rattrapage', 'exam', 'rat']

    def _resolve_col(candidates):
        for c in candidates:
            if c in headers:
                return c
        return None

    if type_note:
        if type_note == 'EXAM':
            return {'EXAM': _resolve_col(EXAM_ALIASES) or 'note_exam'}
        return {type_note: f'note_{type_note.lower()}'}
    if type_session == 'rattrapage':
        return {'EXAM': _resolve_col(EXAM_ALIASES) or 'note_exam'}
    return {'CC': 'note_cc', 'TP': 'note_tp', 'EXAM': 'note_exam'}


def _index_inscriptions(element_id):
    """matricule -> InscriptionElement pour cet EM (1 requête)."""
    from apps.inscriptions.models import InscriptionElement
    return {
        ie.inscription_ped.inscription_admin.etudiant.matricule: ie
        for ie in InscriptionElement.objects.filter(em_id=element_id).select_related(
            'inscription_ped__inscription_admin__etudiant'
        )
    }


def _importer_lignes(ws, headers, col_map, ie_index, session, user, eligible_ids=None):
    """Parcourt les lignes xlsx, décide create/update en mémoire puis écrit en
    masse (bulk_create + bulk_update). Retourne (created, updated, errors).

    eligible_ids : None (session normale → aucun filtre) ou un set d'IE éligibles
    (session rattrapage → seuls ces IE sont saisissables ; les autres sont ignorés
    avec une ligne d'erreur). Voir note_lecture.eligible_rattrapage_ie_ids."""
    from decimal import Decimal, InvalidOperation
    from django.db import transaction
    from django.utils import timezone
    from apps.evaluations.signals import _local as _signal_local

    def get_val(row, col):
        idx = headers.get(col)
        if idx is None or idx >= len(row) or row[idx] is None:
            return None
        try:
            return Decimal(str(row[idx]).replace(',', '.'))
        except InvalidOperation:
            return None

    created = updated = 0
    errors = []
    # Pré-charge les notes existantes (session, types visés, IE concernés) en 1 requête.
    notes_existantes = {
        (n.inscription_element_id, n.type_note): n
        for n in Note.objects.filter(
            session=session,
            type_note__in=list(col_map.keys()),
            inscription_element__in=list(ie_index.values()),
        )
    }
    a_creer, a_maj, _maj_pks = [], [], set()
    _now = timezone.now()
    # Désactive le recalcul individuel pendant l'import bulk (recalcul global ensuite).
    _signal_local.disabled = True
    try:
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            matricule_cell = headers.get('matricule')
            if matricule_cell is None or matricule_cell >= len(row) or not row[matricule_cell]:
                continue
            matricule = _normalize_matricule(row[matricule_cell])
            if not matricule:
                continue

            ie = ie_index.get(matricule)
            if not ie:
                errors.append({'row': row_idx, 'message': f'Matricule "{matricule}" non trouvé ou non inscrit à cet EM.'})
                continue

            # Garde-fou rattrapage : n'accepter que les IE éligibles (même ensemble
            # que la feuille de saisie). Un étudiant non concerné est IGNORÉ, avec
            # une ligne d'erreur explicite — plus d'écriture SR à tort par import.
            if eligible_ids is not None and ie.id not in eligible_ids:
                errors.append({'row': row_idx,
                               'message': f'Matricule "{matricule}" non concerné par le rattrapage de cet EM — note ignorée.'})
                continue

            for tn, col in col_map.items():
                if col not in headers:
                    continue
                valeur = get_val(row, col)
                if valeur is None:
                    continue  # cellule vide -> on ne touche pas la note existante

                if not (0 <= valeur <= 20):
                    errors.append({'row': row_idx, 'message': f'Note {tn} hors plage [0-20] pour {matricule} : {valeur}.'})
                    continue

                key = (ie.id, tn)
                note = notes_existantes.get(key)
                if note is None:
                    note = Note(
                        inscription_element=ie, session=session, type_note=tn,
                        valeur=valeur, saisie_par=user,
                    )
                    notes_existantes[key] = note
                    a_creer.append(note)
                    created += 1
                else:
                    note.valeur = valeur
                    note.saisie_par = user
                    note.date_modification = _now  # auto_now non déclenché par bulk_update
                    if note.pk is not None and note.pk not in _maj_pks:
                        _maj_pks.add(note.pk)
                        a_maj.append(note)
                        updated += 1

        with transaction.atomic():
            if a_creer:
                Note.objects.bulk_create(a_creer, batch_size=1000)
            if a_maj:
                Note.objects.bulk_update(
                    a_maj, ['valeur', 'saisie_par', 'date_modification'], batch_size=1000,
                )
    except Exception as exc:
        errors.append({'row': 0, 'message': f'Écriture en masse des notes : {exc}'})
    finally:
        _signal_local.disabled = False

    return created, updated, errors


def do_importer(request):
    """POST /api/v1/evaluations/notes/importer/ : import xlsx des notes.

    Colonnes : matricule, note_cc/note_tp/note_exam. Orchestration : validation ->
    chargement xlsx + en-têtes -> mapping colonnes -> index inscriptions -> écriture
    en masse -> recalcul global de la session.
    """
    import openpyxl

    fichier    = request.FILES.get('fichier')
    session_id = request.data.get('session')
    element_id = request.data.get('element')
    type_note  = request.data.get('type_note', '').strip().upper() or None

    if not fichier:
        return Response({'detail': 'Le fichier est requis.'}, status=status.HTTP_400_BAD_REQUEST)
    if fichier.size > 5 * 1024 * 1024:
        return Response({'detail': 'Fichier trop volumineux (max 5 Mo).'}, status=status.HTTP_400_BAD_REQUEST)
    if not session_id or not element_id:
        return Response({'detail': 'session et element sont requis.'}, status=status.HTTP_400_BAD_REQUEST)

    # Scope enseignant : un prof n'importe que pour SES EMs (admin/scolarite/DE : tous).
    if not peut_acceder_em(request.user, element_id):
        return Response({'detail': "Accès non autorisé : import sur un EM que vous n'enseignez pas."},
                        status=status.HTTP_403_FORBIDDEN)

    try:
        session = SessionEvaluation.objects.get(pk=session_id)
    except SessionEvaluation.DoesNotExist:
        return Response({'detail': 'Session introuvable.'}, status=status.HTTP_404_NOT_FOUND)
    if session.est_close:
        return Response({'detail': 'La session est clôturée, aucune saisie autorisée.'}, status=status.HTTP_400_BAD_REQUEST)

    try:
        wb = openpyxl.load_workbook(fichier, read_only=True, data_only=True)
        ws = wb.active
    except Exception as e:
        return Response({'detail': f'Fichier Excel invalide : {e}'}, status=status.HTTP_400_BAD_REQUEST)
    if ws.max_row and ws.max_row > 5000:
        return Response({'detail': 'Trop de lignes (max 5000).'}, status=status.HTTP_400_BAD_REQUEST)

    first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not first_row:
        return Response({'detail': 'Fichier vide.'}, status=status.HTTP_400_BAD_REQUEST)
    headers = {str(c).strip().lower(): i for i, c in enumerate(first_row) if c is not None}
    if 'matricule' not in headers:
        return Response({'detail': 'Colonne "matricule" manquante.'}, status=status.HTTP_400_BAD_REQUEST)

    col_map = _resoudre_col_map(headers, type_note, session.type_session)
    ie_index = _index_inscriptions(element_id)

    # Garde-fou : en RATTRAPAGE, n'accepter une note que pour les IE éligibles
    # (obligation OU dette non validée) — MÊME ensemble que la feuille de saisie.
    # En session normale : aucun filtre (tous les inscrits à l'EM sont saisissables).
    eligible_ids = None
    if session.type_session == 'rattrapage':
        from apps.evaluations.services.note_lecture import eligible_rattrapage_ie_ids
        eligible_ids = eligible_rattrapage_ie_ids(element_id, session)

    created, updated, errors = _importer_lignes(
        ws, headers, col_map, ie_index, session, request.user, eligible_ids,
    )

    # Recalcul global de la session après l'import bulk (signal désactivé pendant).
    try:
        svc = NoteCalculService(session)
        svc.calculer_tous_elements_session()
        svc.calculer_tous_modules_session()
        svc.calculer_tous_semestres_session()
    except Exception as exc:
        errors.append({'row': 0, 'message': f'Recalcul global apres import : {exc}'})

    return Response({'created': created, 'updated': updated, 'errors': errors})
def do_saisir_bulk(request):
    """
    POST { session: ID, rows: [{inscription_element: ID, cc, tp, exam}] }
    Upsert les notes CC/TP/EXAM. Valeur vide ('' ou null) = PAS de note → la note
    existante est SUPPRIMÉE (décision : double vérification humaine, risque
    d'effacement accidentel jugé négligeable). Valeur invalide/hors [0,20] = ignorée.
    """
    from decimal import Decimal, InvalidOperation
    from django.db import transaction
    from apps.inscriptions.models import InscriptionElement

    session_id = request.data.get('session')
    rows       = request.data.get('rows', [])
    if not session_id:
        return Response({'detail': 'session requis.'}, status=status.HTTP_400_BAD_REQUEST)
    try:
        session = SessionEvaluation.objects.get(pk=session_id)
    except SessionEvaluation.DoesNotExist:
        return Response({'detail': 'Session introuvable.'}, status=status.HTTP_404_NOT_FOUND)
    if session.est_close:
        return Response({'detail': 'Session clôturée — saisie impossible.'}, status=status.HTTP_400_BAD_REQUEST)

    # Scope enseignant : verifier que TOUS les EMs vises par les lignes sont
    # autorises pour cet utilisateur (un prof ne saisit que pour SES EMs).
    ie_ids = [r.get('inscription_element') for r in rows if r.get('inscription_element')]
    em_ids = set(
        InscriptionElement.objects.filter(id__in=ie_ids)
        .values_list('em_id', flat=True)
    )
    for em_id in em_ids:
        if not peut_acceder_em(request.user, em_id):
            return Response(
                {'detail': 'Accès non autorisé : saisie sur un EM que vous n\'enseignez pas.'},
                status=status.HTTP_403_FORBIDDEN,
            )

    created = updated = deleted = 0
    with transaction.atomic():
        for row in rows:
            ie_id = row.get('inscription_element')
            if not ie_id:
                continue
            for type_note, val_key in (('CC', 'cc'), ('TP', 'tp'), ('EXAM', 'exam')):
                raw = row.get(val_key)
                if raw is None or raw == '':
                    # Champ vidé = pas de note → SUPPRIME la note existante (uniquement
                    # pour ce (inscription_element, CETTE session, type_note) ; n'affecte
                    # pas la session normale lors d'une saisie de rattrapage).
                    nb, _ = Note.objects.filter(
                        inscription_element_id=ie_id,
                        session=session,
                        type_note=type_note,
                    ).delete()
                    if nb:
                        deleted += 1
                    continue
                try:
                    valeur = Decimal(str(raw).replace(',', '.'))
                except InvalidOperation:
                    continue
                if not (0 <= valeur <= 20):
                    continue
                note, was_created = Note.objects.update_or_create(
                    inscription_element_id=ie_id,
                    session=session,
                    type_note=type_note,
                    defaults={'valeur': valeur, 'saisie_par': request.user},
                )
                if was_created:
                    created += 1
                else:
                    updated += 1

    return Response({'created': created, 'updated': updated, 'deleted': deleted})
